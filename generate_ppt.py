"""
Generate Professional PPTX Presentation for Clinic Management System Capstone Review
"""
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen format
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    NAVY = RGBColor(9, 20, 36)
    BLUE = RGBColor(2, 132, 199)
    TEAL = RGBColor(13, 148, 136)
    LIGHT_BG = RGBColor(248, 250, 252)
    DARK_CARD = RGBColor(15, 30, 54)
    TEXT_MUTED = RGBColor(100, 116, 139)
    WHITE = RGBColor(255, 255, 255)
    GREEN = RGBColor(5, 150, 105)

    blank_layout = prs.slide_layouts[6]
    screenshots_dir = os.path.join(os.path.dirname(__file__), 'screenshots')

    def add_header(slide, eyebrow_text, title_text, dark=False):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = eyebrow_text.upper()
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = BLUE if not dark else RGBColor(56, 189, 248)

        p2 = tf.add_paragraph()
        p2.text = title_text
        p2.font.size = Pt(24)
        p2.font.bold = True
        p2.font.color.rgb = NAVY if not dark else WHITE
        p2.space_before = Pt(4)

    # -------------------------------------------------------------------------
    # SLIDE 1: Title Slide (Dark Theme)
    # -------------------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.color.rgb = NAVY

    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11), Inches(3.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "DBMS CAPSTONE PROJECT PRESENTATION"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)

    p = tf1.add_paragraph()
    p.text = "MediCare+ Clinic Management System"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_before = Pt(12)

    p = tf1.add_paragraph()
    p.text = "An Enterprise Healthcare Relational Platform: Backend SQL to Frontend Integration"
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(148, 163, 184)
    p.space_before = Pt(8)

    p = tf1.add_paragraph()
    p.text = "• Database Engine: MySQL 8.0+ Relational Schema (3NF, Triggers, Views, Stored Procedures)\n• Application Layer: Python Flask Architecture & Session Management\n• Frontend: Responsive Healthcare Design System (HTML5, CSS3, JavaScript)"
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(203, 213, 225)
    p.space_before = Pt(20)

    # -------------------------------------------------------------------------
    # SLIDE 2: Step-by-Step Development Procedure (Backend to Frontend)
    # -------------------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Architecture & Workflow", "Step-by-Step Development Procedure: Backend to Frontend")

    steps = [
        ("Step 1: Database Schema & 3NF Normalization", "Designed 7 core entities with primary keys, auto-increments, check constraints, and cascading foreign keys in MySQL."),
        ("Step 2: Advanced SQL Programmability", "Created Stored Procedures for atomic booking/consultation, Triggers for automatic 5% GST billing, and Analytical Views for real-time reporting."),
        ("Step 3: Database Verification & Connector Layer", "Built automated initialization script (init_db.py) and a fail-safe dual-engine database adapter in Flask with connection pooling."),
        ("Step 4: Flask Backend & Authentication API", "Developed 18 RESTful endpoints, secure password hashing, and automatic multi-role recognition (Patient, Doctor, Admin)."),
        ("Step 5: Frontend Design System & Clinical UI", "Implemented modern CSS custom properties, Google Fonts, responsive medical cards, printable invoice stylesheet, and live table search.")
    ]

    top_pos = 1.8
    for i, (title, desc) in enumerate(steps, 1):
        box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(top_pos), Inches(11.3), Inches(0.85))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = RGBColor(226, 232, 240)

        # Step badge
        badge = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.15), Inches(top_pos + 0.15), Inches(0.9), Inches(0.55))
        badge.fill.solid()
        badge.fill.fore_color.rgb = BLUE
        badge.line.color.rgb = BLUE
        btf = badge.text_frame
        bp = btf.paragraphs[0]
        bp.text = f"0{i}"
        bp.font.size = Pt(14)
        bp.font.bold = True
        bp.font.color.rgb = WHITE
        bp.alignment = PP_ALIGN.CENTER

        # Text inside step
        txt = s2.shapes.add_textbox(Inches(2.2), Inches(top_pos + 0.08), Inches(10.0), Inches(0.7))
        ttf = txt.text_frame
        ttf.word_wrap = True
        tp1 = ttf.paragraphs[0]
        tp1.text = title
        tp1.font.size = Pt(13.5)
        tp1.font.bold = True
        tp1.font.color.rgb = NAVY

        tp2 = ttf.add_paragraph()
        tp2.text = desc
        tp2.font.size = Pt(11)
        tp2.font.color.rgb = TEXT_MUTED

        top_pos += 1.02

    # -------------------------------------------------------------------------
    # SLIDE 3: Phase 1 — Database Architecture & Relational Tables
    # -------------------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Phase 1: Backend Database Engineering", "Relational Schema Design & 3NF Entity Normalization")

    entities = [
        ("users", "user_id (PK), username (UQ), password, role, is_active, created_at", "Role-based authentication table for patients, doctors, and administrators."),
        ("department", "DepartmentID (PK), DepartmentName (UQ), Description, HeadOfDepartment", "Clinical specialties (Cardiology, Dermatology, Pediatrics, General Medicine)."),
        ("patient", "PatientID (PK), PatientName, DOB, Gender, Phone (UQ), BloodGroup, Address, UserID (FK)", "Medical demographics, blood group cataloging, emergency contact."),
        ("doctor", "DoctorID (PK), DoctorName, Specialization, Phone, DepartmentID (FK), ConsultationFee, RoomNumber, UserID (FK)", "Medical staff directory, consultation charges, assigned room numbers."),
        ("appointment", "AppointmentID (PK), PatientID (FK), DoctorID (FK), AppointmentDate, Reason, Symptoms, Diagnosis, Status", "Core clinical consultation ledger tracking scheduled, confirmed, completed visits."),
        ("prescription", "PrescriptionID (PK), AppointmentID (FK), Medicine, Dosage, Duration, Instructions", "Itemized electronic pharmacy orders issued by licensed physicians."),
        ("bill", "BillID (PK), AppointmentID (FK, UQ), Amount, Tax (5%), TotalAmount, PaymentStatus, PaymentMethod, PaidAt", "Financial accounting ledger with GST computation and settlement timestamps.")
    ]

    top_pos = 1.7
    for name, schema, purpose in entities:
        box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(top_pos), Inches(11.3), Inches(0.68))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = RGBColor(226, 232, 240)

        txt = s3.shapes.add_textbox(Inches(1.2), Inches(top_pos + 0.05), Inches(10.9), Inches(0.6))
        ttf = txt.text_frame
        ttf.word_wrap = True

        p1 = ttf.paragraphs[0]
        p1.text = f"{name.upper()} TABLE  —  {schema}"
        p1.font.size = Pt(11.5)
        p1.font.bold = True
        p1.font.color.rgb = BLUE

        p2 = ttf.add_paragraph()
        p2.text = f"Purpose: {purpose}"
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_MUTED

        top_pos += 0.77

    # -------------------------------------------------------------------------
    # SLIDE 4: Phase 2 — Advanced SQL Programmability (Procedures, Triggers, Views)
    # -------------------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Phase 2: Database Programmability", "Advanced SQL: Stored Procedures, Triggers, and Analytical Views")

    # Column 1: Stored Procedures & Triggers
    c1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.8), Inches(5.4), Inches(5.1))
    c1.fill.solid()
    c1.fill.fore_color.rgb = WHITE
    c1.line.color.rgb = RGBColor(226, 232, 240)

    txt1 = s4.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(5.0), Inches(4.7))
    tf = txt1.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "⚡ Stored Procedures & Triggers"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY

    items_c1 = [
        ("sp_book_appointment", "Validates doctor availability, verifies consultation fee, inserts appointment and generates initial bill atomically."),
        ("sp_complete_consultation", "Atomically updates diagnosis, clinical notes, marks appointment as 'Completed' and inserts medicine details."),
        ("sp_pay_bill", "Updates invoice status to 'Paid', records payment mode (UPI/Card/Cash) and logs settlement timestamp."),
        ("trg_appointment_auto_bill", "Trigger that automatically calculates doctor fee + 5% GST and creates a pending bill record on appointment insertion."),
        ("trg_prevent_past_appointment", "Integrity trigger that prevents booking consultation dates in the past.")
    ]
    for title, desc in items_c1:
        p1 = tf.add_paragraph()
        p1.text = f"• {title}"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = BLUE
        p1.space_before = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_MUTED

    # Column 2: Analytical SQL Views
    c2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.4), Inches(5.1))
    c2.fill.solid()
    c2.fill.fore_color.rgb = WHITE
    c2.line.color.rgb = RGBColor(226, 232, 240)

    txt2 = s4.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.0), Inches(4.7))
    tf2 = txt2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "📊 Relational SQL Views"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY

    items_c2 = [
        ("patient_appointment_report", "Multi-table relational JOIN uniting Patient, Doctor, Department, Appointment status, and Invoice information."),
        ("v_doctor_schedules", "Aggregates doctor caseloads, upcoming bookings, and completed patient consultations per physician."),
        ("v_billing_summary", "Departmental financial breakdown: Total Billed, Total Collected Revenue, and Outstanding Receivables."),
        ("v_clinic_stats", "Instant executive summary view tracking total active patients, doctors, scheduled visits, and hospital receipts."),
        ("Advanced Queries (queries.sql)", "Features 4-way relational JOINs, GROUP BY with HAVING, and Window Functions (ROW_NUMBER, DENSE_RANK).")
    ]
    for title, desc in items_c2:
        p1 = tf2.add_paragraph()
        p1.text = f"• {title}"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = TEAL
        p1.space_before = Pt(8)

        p2 = tf2.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------------------
    # SLIDE 5: Phase 3 — Backend Web Layer & Security (Flask + MySQL Connector)
    # -------------------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Phase 3: Application Logic Layer", "Python Flask Backend, Role Security & Business Workflows")

    modules = [
        ("Seamless Dual Database Adapter", "Connects directly to MySQL 8.0+ server (clinic_management) via mysql.connector, with automated schema loader and local fail-safe mirror."),
        ("Automated Multi-Role Authentication", "Eliminated cumbersome role dropdowns. The system securely looks up username/password and automatically recognizes whether the user is a Patient, Doctor, or Admin."),
        ("Patient Self-Enrollment & Doctor Onboarding", "Full registration API (/register) enabling new patients to register with blood group and contact info, and medical doctors to join departments."),
        ("Clinical e-Prescription Workflow", "Doctors can prescribe medicine directly from appointment rows with dosage, duration, and instructions, updating diagnosis in one transaction."),
        ("Billing & Instant Invoice Generation", "Provisional invoice generated upon appointment creation; users can settle invoices online (UPI/Card/Cash) and print clinical receipts.")
    ]

    top_pos = 1.8
    for title, desc in modules:
        box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(top_pos), Inches(11.3), Inches(0.88))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = RGBColor(226, 232, 240)

        txt = s5.shapes.add_textbox(Inches(1.3), Inches(top_pos + 0.1), Inches(10.8), Inches(0.7))
        ttf = txt.text_frame
        ttf.word_wrap = True

        p1 = ttf.paragraphs[0]
        p1.text = f"⚙️ {title}"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = NAVY

        p2 = ttf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED

        top_pos += 1.02

    # -------------------------------------------------------------------------
    # SLIDE 6: Implementation Screenshot — Home & Clinical Services
    # -------------------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Implementation Progress: Frontend", "Landing Page: Medical OPD Care, 24/7 Helpline & Live Metrics")

    # Add screenshot image if available
    img_path = os.path.join(screenshots_dir, 'Screenshot 2026-10-04 214644.png')
    if os.path.exists(img_path):
        s6.shapes.add_picture(img_path, Inches(1.0), Inches(1.8), Inches(6.5))

    desc_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.8), Inches(4.5), Inches(5.1))
    desc_box.fill.solid()
    desc_box.fill.fore_color.rgb = WHITE
    desc_box.line.color.rgb = RGBColor(226, 232, 240)

    txt = s6.shapes.add_textbox(Inches(8.0), Inches(2.0), Inches(4.1), Inches(4.7))
    tf = txt.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Clinical Landing Page Features"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY

    bullets = [
        ("Live OPD Emergency Card", "Displays clinic consultation hours (Mon-Sat 8AM-8PM), 24/7 emergency helpline, and instant appointment booking CTA."),
        ("Clickable KPI Stat Cards", "All metric counters (Total Patients, Specialists, Consultations, Departments) are interactive clickable links leading directly to respective directories."),
        ("Doctor Showcase Grid", "Highlights top doctors, specializations, consulting rooms, and consultation charges."),
        ("Modern Healthcare Aesthetics", "Styled with Plus Jakarta Sans typography, frosted glass header, and responsive layout.")
    ]
    for b_title, b_desc in bullets:
        p1 = tf.add_paragraph()
        p1.text = f"• {b_title}: {b_desc}"
        p1.font.size = Pt(10.5)
        p1.font.color.rgb = TEXT_MUTED
        p1.space_before = Pt(8)

    # -------------------------------------------------------------------------
    # SLIDE 7: Implementation Screenshot — Interactive Dashboard
    # -------------------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Implementation Progress: Dashboard", "Role-Tailored Dashboard with Clickable Metric Cards & Schedule")

    img_path = os.path.join(screenshots_dir, 'Screenshot 2026-10-04 211405.png')
    if os.path.exists(img_path):
        s7.shapes.add_picture(img_path, Inches(1.0), Inches(1.8), Inches(6.5))

    desc_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.8), Inches(4.5), Inches(5.1))
    desc_box.fill.solid()
    desc_box.fill.fore_color.rgb = WHITE
    desc_box.line.color.rgb = RGBColor(226, 232, 240)

    txt = s7.shapes.add_textbox(Inches(8.0), Inches(2.0), Inches(4.1), Inches(4.7))
    tf = txt.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Dashboard Architecture"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY

    bullets = [
        ("Role-Adaptive Views", "Adapts interface dynamically based on role: Patient sees their visits/bills; Doctor sees scheduled patients; Admin monitors entire hospital operations."),
        ("Interactive Navigation", "Clicking any KPI card (Patients, Doctors, Appointments, Pending Bills) immediately navigates to filtered records."),
        ("Cashflow Audit Banner", "Real-time summary of collected revenue and pending receivables powered by SQL SUM() aggregate calculations."),
        ("Recent Consultation Ledger", "Lists upcoming visits with direct status badges and invoice links.")
    ]
    for b_title, b_desc in bullets:
        p1 = tf.add_paragraph()
        p1.text = f"• {b_title}: {b_desc}"
        p1.font.size = Pt(10.5)
        p1.font.color.rgb = TEXT_MUTED
        p1.space_before = Pt(8)

    # -------------------------------------------------------------------------
    # SLIDE 8: Implementation Screenshot — Appointments & Interactive Billing
    # -------------------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Implementation Progress: Clinical Ledger", "Appointment Management with Clickable Invoice & Status Filter")

    img_path = os.path.join(screenshots_dir, 'Screenshot 2026-10-04 211644.png')
    if os.path.exists(img_path):
        s8.shapes.add_picture(img_path, Inches(1.0), Inches(1.8), Inches(6.5))

    desc_box = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.8), Inches(4.5), Inches(5.1))
    desc_box.fill.solid()
    desc_box.fill.fore_color.rgb = WHITE
    desc_box.line.color.rgb = RGBColor(226, 232, 240)

    txt = s8.shapes.add_textbox(Inches(8.0), Inches(2.0), Inches(4.1), Inches(4.7))
    tf = txt.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Appointment & Bill Workflow"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY

    bullets = [
        ("Status Tabs & Live Search", "Filter visits by Scheduled, Confirmed, Completed, Cancelled, or use instant client-side table search."),
        ("Clickable Bill Status", "Clicking on 'Paid' or 'Pending' bill badges instantly opens the itemized clinical receipt and invoice view."),
        ("Clean Datetime Formatting", "Human-readable appointment schedule display without raw ISO separators."),
        ("Doctor Prescription Action", "Doctors can update diagnosis and write electronic prescriptions directly from their appointment table rows.")
    ]
    for b_title, b_desc in bullets:
        p1 = tf.add_paragraph()
        p1.text = f"• {b_title}: {b_desc}"
        p1.font.size = Pt(10.5)
        p1.font.color.rgb = TEXT_MUTED
        p1.space_before = Pt(8)

    # -------------------------------------------------------------------------
    # SLIDE 9: Implementation Screenshot — Electronic Prescriptions
    # -------------------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Implementation Progress: Pharmacy", "Digital Prescription Cards with Attending Doctor & Patient Attribution")

    img_path = os.path.join(screenshots_dir, 'Screenshot 2026-10-04 214430.png')
    if os.path.exists(img_path):
        s9.shapes.add_picture(img_path, Inches(1.0), Inches(1.8), Inches(6.5))

    desc_box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.8), Inches(4.5), Inches(5.1))
    desc_box.fill.solid()
    desc_box.fill.fore_color.rgb = WHITE
    desc_box.line.color.rgb = RGBColor(226, 232, 240)

    txt = s9.shapes.add_textbox(Inches(8.0), Inches(2.0), Inches(4.1), Inches(4.7))
    tf = txt.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Digital Prescription Architecture"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY

    bullets = [
        ("Medication Details", "Displays medicine name, clinical dosage (e.g. 25 mg), duration period, and clear directions (e.g. After meals)."),
        ("Clinical Diagnosis", "Displays verified medical diagnosis linked to the specific consultation visit."),
        ("Prescribing Doctor Name", "Accredits attending physician, clinical specialization, and hospital department (e.g., Dr. Priya Sharma • Cardiology)."),
        ("Pharmacy Verification", "Provides authentic e-prescription record ready for pharmacy dispensing and record-keeping.")
    ]
    for b_title, b_desc in bullets:
        p1 = tf.add_paragraph()
        p1.text = f"• {b_title}: {b_desc}"
        p1.font.size = Pt(10.5)
        p1.font.color.rgb = TEXT_MUTED
        p1.space_before = Pt(8)

    # -------------------------------------------------------------------------
    # SLIDE 10: Conclusion & Outcomes
    # -------------------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    bg10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg10.fill.solid()
    bg10.fill.fore_color.rgb = NAVY
    bg10.line.color.rgb = NAVY

    add_header(s10, "Summary & Future Roadmap", "Project Conclusion & Technical Accomplishments", dark=True)

    # 3 Summary Cards
    card_data = [
        ("Relational Integrity", "Implemented robust 3NF normalized schema with MySQL triggers, auto-billing, stored procedures, and analytical views eliminating data redundancy."),
        ("End-to-End Integration", "Successfully connected MySQL database backend with Python Flask routes and an ultra-crisp responsive frontend with real-time feedback."),
        ("Seamless Clinical UX", "Delivered automatic role-based login, clickable navigation, printable GST receipts, and verified doctor prescriptions meeting industry standards.")
    ]

    left_pos = 1.0
    for title, desc in card_data:
        box = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_pos), Inches(2.2), Inches(3.6), Inches(4.2))
        box.fill.solid()
        box.fill.fore_color.rgb = DARK_CARD
        box.line.color.rgb = RGBColor(26, 47, 78)

        txt = s10.shapes.add_textbox(Inches(left_pos + 0.2), Inches(2.5), Inches(3.2), Inches(3.6))
        ttf = txt.text_frame
        ttf.word_wrap = True

        p1 = ttf.paragraphs[0]
        p1.text = f"✓ {title}"
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(56, 189, 248)

        p2 = ttf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = RGBColor(203, 213, 225)
        p2.space_before = Pt(14)

        left_pos += 3.85

    output_path = os.path.join(os.path.dirname(__file__), 'MediCare_Clinic_Management_Project_Presentation.pptx')
    prs.save(output_path)
    print(f"[+] PPTX successfully generated at: {output_path}")
    return output_path

if __name__ == '__main__':
    create_presentation()
